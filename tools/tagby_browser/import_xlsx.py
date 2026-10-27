#!/usr/bin/env python3
"""Stage an already-downloaded, authorized TAGby similar-accounts workbook locally.
No login, search, network, payment or production DB operation is performed.
"""
import argparse,json,re,os,tempfile
from pathlib import Path
from decimal import Decimal,InvalidOperation
from urllib.parse import urlparse
from datetime import datetime,timezone
HEADERS=['유사도','SNS ID','SNS URL','평균 좋아요','평균 댓글','팔로워']
HANDLE=re.compile(r'^[a-z0-9._]{1,30}$')
def metric(value):
 if value is None or isinstance(value,bool):return None
 s=str(value).strip().replace(',','');m=re.fullmatch(r'(\d+(?:\.\d+)?)(만|천)?',s)
 if not m:return None
 try:n=Decimal(m[1])*({'만':10000,'천':1000}.get(m[2],1));return int(n) if n>=0 and n==int(n) else None
 except (ValueError,InvalidOperation):return None

def from_workbook(file,seed=''):
 from openpyxl import load_workbook
 if seed and not HANDLE.fullmatch(seed):raise ValueError('기준 아이디 형식 오류')
 wb=load_workbook(file,read_only=True,data_only=False)
 try:
  if len(wb.worksheets)!=1:raise ValueError('검증된 한 개 시트 형식만 지원합니다.')
  it=wb.active.iter_rows(values_only=False);head=next(it,None)
  if not head or [str(c.value or '').strip() for c in head[:6]]!=HEADERS:raise ValueError('검증된 TAGby 유사 계정 열 형식이 아닙니다.')
  result=[];seen=set()
  for row in it:
   if all(c.value is None for c in row):continue
   if len(result)>=30:raise ValueError('검증 단계는 결과 30행 이하만 허용합니다.')
   if any(c.data_type=='f' for c in row[:6]):raise ValueError('수식 셀이 포함돼 자동으로 읽지 않습니다.')
   v=[c.value for c in row[:6]];handle=str(v[1] or '').strip().lstrip('@').lower()
   if not HANDLE.fullmatch(handle):raise ValueError('SNS ID 확인 필요')
   if handle in seen:continue
   seen.add(handle);url=str(v[2] or '').strip();profile=''
   if url:
    u=urlparse(url);segments=[x for x in u.path.split('/') if x]
    if u.scheme!='https' or u.hostname not in ('instagram.com','www.instagram.com') or len(segments)!=1 or segments[0].lower()!=handle:raise ValueError('아이디와 Instagram URL 불일치')
    profile='https://'+u.hostname+'/'+segments[0]+'/'
   result.append({'handle':handle,'source_id':'tagby:'+handle,'profile_url':profile,'followers':metric(v[5]),'followers_text':str(v[5] or ''),'category':'','similarity':str(v[0] or '')[:100],'average_likes':metric(v[3]),'average_comments':metric(v[4]),'source_url':'https://tagby.io/instant/recommend/search','source_ref':seed,'checked':datetime.now(timezone.utc).isoformat(),'observation_kind':'file_observed','source_position':len(result)+1,'email_status':'아직 조회하지 않음','assessment':'review','reasons':['태그바이 유사 계정 결과 · 유사도는 공동구매 적합성을 의미하지 않음']})
  return result
 finally:wb.close()
def stage(rows,file):
 if file.exists():
  state=json.loads(file.read_text(encoding='utf8'))
  if not isinstance(state.get('rows'),list) or not isinstance(state.get('scrollY'),(int,float)):raise ValueError('기존 진행 기록 오류')
 else:state={'rows':[],'scrollY':0}
 keys={(r['handle'],r.get('source_url',''),r.get('source_ref','')) for r in state['rows']};added=0
 for row in rows:
  key=(row['handle'],row['source_url'],row['source_ref'])
  if key not in keys:state['rows'].append(row);keys.add(key);added+=1
 file.parent.mkdir(mode=0o700,parents=True,exist_ok=True)
 fd,tmp=tempfile.mkstemp(prefix='.tagby-',dir=file.parent)
 try:
  if hasattr(os,'fchmod'):os.fchmod(fd,0o600)
  with os.fdopen(fd,'w',encoding='utf8') as out:json.dump(state,out,ensure_ascii=False,indent=2)
  os.replace(tmp,file)
 finally:
  if os.path.exists(tmp):os.unlink(tmp)
 return added
if __name__=='__main__':
 p=argparse.ArgumentParser();p.add_argument('xlsx',type=Path);p.add_argument('--seed',default='');p.add_argument('--checkpoint',type=Path,default=Path.home()/'.sanji-tagby'/'checkpoint.json');a=p.parse_args()
 rows=from_workbook(a.xlsx,a.seed.lstrip('@').lower());print(json.dumps({'read':len(rows),'newEvidence':stage(rows,a.checkpoint),'destination':'local checkpoint only','productionChanged':False},ensure_ascii=False))
